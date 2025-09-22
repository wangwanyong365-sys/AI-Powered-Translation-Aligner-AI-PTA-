import os
import json
import time
import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog

import openai
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

from app_utils import log_error, save_settings, translate_single_paragraph

RESUME_PE_FILE = "resume_post_edit.json"


class PostEditingWindow(tk.Toplevel):
    def __init__(self, parent):
        super().__init__(parent)
        self.withdraw()
        self.parent = parent
        self.title("Post-editing Tool")
        self.geometry("700x600")
        self.minsize(600, 500)
        self.columnconfigure(0, weight=1)
        self.rowconfigure(0, weight=1)

        self.selected_files = []
        self.stop_requested = threading.Event()
        self.is_processing = False
        self.resume_data = None
        self.timer_id = None
    
        self._setup_style()
        self._setup_ui()
        self._post_ui_setup()
        self.protocol("WM_DELETE_WINDOW", self._on_closing)
        self.transient(parent)
        self.grab_set()
        self.after(100, self._check_for_resume_task)
        self.deiconify()
    
    def _on_closing(self):
        if self.is_processing:
            if not messagebox.askyesno("Confirm Exit", "A post-editing task is in progress. Exiting now will stop it. Are you sure?", parent=self):
                return
            self.stop_requested.set()
        self.destroy()
    
    def _setup_style(self):
        self.style = ttk.Style(self)
        self.style.theme_use("clam")
        
        default_font = ("Segoe UI", 10)
        self.style.configure("TLabel", font=default_font)
        self.style.configure("TButton", font=default_font, padding=5)
        self.style.configure("TEntry", font=default_font)
        
        self.style.configure("TLabelframe.Label", font=("Segoe UI", 12, "bold"))
        
        self.style.configure("Accent.TButton", font=("Segoe UI", 10, "bold"), padding=8)
        self.style.map("Accent.TButton",
                       background=[("active", "#0078D7")],
                       foreground=[("active", "white")])
    
    def _setup_ui(self):
        main_frame = ttk.Frame(self, padding="15")
        main_frame.grid(row=0, column=0, sticky="nsew")
        main_frame.rowconfigure(1, weight=1)
        main_frame.columnconfigure(0, weight=1)
    
        file_frame = ttk.LabelFrame(main_frame, text="File Selection", padding="10")
        file_frame.grid(row=0, column=0, sticky="ew", pady=(0, 10))
        file_frame.columnconfigure(0, weight=1)
        
        self.file_listbox = tk.Listbox(file_frame, height=4, font=("Segoe UI", 10))
        self.file_listbox.grid(row=0, column=0, sticky="ew")
        scrollbar = ttk.Scrollbar(file_frame, orient=tk.VERTICAL, command=self.file_listbox.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        self.file_listbox.config(yscrollcommand=scrollbar.set)
        browse_button = ttk.Button(file_frame, text="Select XLSX Files...", command=self._browse_files)
        browse_button.grid(row=1, column=0, columnspan=2, pady=(10, 0), sticky="e")
    
        prompt_frame = ttk.LabelFrame(main_frame, text="Post-editing Prompt", padding="10")
        prompt_frame.grid(row=1, column=0, sticky="nsew", pady=5)
        prompt_frame.columnconfigure(0, weight=1)
        prompt_frame.rowconfigure(1, weight=1)
        
        prompt_selection_frame = ttk.Frame(prompt_frame)
        prompt_selection_frame.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 5))
        prompt_selection_frame.columnconfigure(1, weight=1)
        
        ttk.Label(prompt_selection_frame, text="Select Prompt:").grid(row=0, column=0, padx=(0, 5))
        self.prompt_var = tk.StringVar()
        self.prompt_combo = ttk.Combobox(prompt_selection_frame, textvariable=self.prompt_var, state="readonly")
        self.prompt_combo.grid(row=0, column=1, sticky="ew")
        self.prompt_combo.bind("<<ComboboxSelected>>", self._on_prompt_select)
        
        btn_frame = ttk.Frame(prompt_selection_frame)
        btn_frame.grid(row=0, column=2, padx=(10, 0))
        ttk.Button(btn_frame, text="Add", command=self._add_prompt, width=6).pack(side=tk.LEFT, padx=2)
        ttk.Button(btn_frame, text="Save", command=self._save_current_prompt, width=6).pack(side=tk.LEFT, padx=2)
        ttk.Button(btn_frame, text="Delete", command=self._delete_prompt, width=6).pack(side=tk.LEFT, padx=2)
        
        self.prompt_text = tk.Text(prompt_frame, wrap=tk.WORD, height=8, font=("Segoe UI", 10))
        self.prompt_text.grid(row=1, column=0, sticky="nsew")
        prompt_scrollbar = ttk.Scrollbar(prompt_frame, orient=tk.VERTICAL, command=self.prompt_text.yview)
        prompt_scrollbar.grid(row=1, column=1, sticky="ns")
        self.prompt_text.config(yscrollcommand=prompt_scrollbar.set)
    
        self.process_button = ttk.Button(main_frame, text="Start Post-editing", command=self._start_post_editing, style="Accent.TButton")
        self.process_button.grid(row=2, column=0, pady=10, sticky="ew")
    
        status_bar = ttk.Frame(self)
        status_bar.grid(row=1, column=0, sticky="ew")
        self.status_label = ttk.Label(status_bar, text="Ready", relief=tk.SUNKEN, anchor=tk.W, padding=5)
        self.status_label.pack(side=tk.LEFT, expand=True, fill=tk.X)
        self.timer_label = ttk.Label(status_bar, text="", relief=tk.SUNKEN, anchor=tk.E, padding=5)
        self.timer_label.pack(side=tk.RIGHT)
        self._update_status("Ready", "gray")
    
    def _post_ui_setup(self):
        self._update_prompt_combo()
        if self.parent.settings['post_editing_prompts']:
            first_prompt_name = list(self.parent.settings['post_editing_prompts'].keys())[0]
            self.prompt_var.set(first_prompt_name)
            self._on_prompt_select()
    
    def _browse_files(self):
        files = filedialog.askopenfilenames(
            title="Select Excel files to post-edit",
            filetypes=(("Excel files", "*.xlsx"), ("All files", "*.*")),
            parent=self
        )
        if files:
            self.selected_files = list(files)
            self.file_listbox.delete(0, tk.END)
            for file in self.selected_files:
                self.file_listbox.insert(tk.END, os.path.basename(file))
            self._update_status(f"Selected {len(self.selected_files)} files", "blue")
    
    def _update_status(self, text, color):
        self.status_label.config(text=text, foreground=color)
        self.update_idletasks()
    
    def _update_timer(self, start_time):
        elapsed = time.time() - start_time
        mins, secs = divmod(elapsed, 60)
        self.timer_label.config(text=f"{int(mins):02d}:{secs:04.1f}")
        self.timer_id = self.after(100, self._update_timer, start_time)
    
    def _cancel_timer(self):
        if self.timer_id:
            self.after_cancel(self.timer_id)
            self.timer_id = None
        self.timer_label.config(text="")
    
    def _update_prompt_combo(self):
        self.prompt_combo['values'] = list(self.parent.settings['post_editing_prompts'].keys())
    
    def _on_prompt_select(self, event=None):
        name = self.prompt_var.get()
        if name in self.parent.settings['post_editing_prompts']:
            self.prompt_text.delete("1.0", tk.END)
            self.prompt_text.insert("1.0", self.parent.settings['post_editing_prompts'][name])
    
    def _add_prompt(self):
        name = simpledialog.askstring("New Prompt", "Enter name for the new Post-editing Prompt:", parent=self)
        if name and name.strip():
            name = name.strip()
            if name in self.parent.settings['post_editing_prompts']:
                messagebox.showerror("Error", "This name already exists!", parent=self)
                return
            self.parent.settings['post_editing_prompts'][name] = "Enter instructions, using {source} and {target} placeholders."
            save_settings(self.parent.settings)
            self._update_prompt_combo()
            self.prompt_var.set(name)
            self._on_prompt_select()
    
    def _save_current_prompt(self):
        name = self.prompt_var.get()
        if not name: return messagebox.showerror("Error", "No prompt selected to save.", parent=self)
        content = self.prompt_text.get("1.0", tk.END).strip()
        if "{source}" not in content or "{target}" not in content:
            if not messagebox.askyesno("Warning", "The placeholders {source} and {target} were not found. This may cause errors.\nDo you still want to save?", parent=self):
                return
        self.parent.settings['post_editing_prompts'][name] = content
        save_settings(self.parent.settings)
        messagebox.showinfo("Success", f"Prompt '{name}' has been saved.", parent=self)
    
    def _delete_prompt(self):
        name = self.prompt_var.get()
        if not name: return messagebox.showerror("Error", "Please select a prompt to delete.", parent=self)
        if len(self.parent.settings['post_editing_prompts']) <= 1:
            return messagebox.showwarning("Warning", "You cannot delete the last prompt.", parent=self)
        if messagebox.askyesno("Confirm Deletion", f"Are you sure you want to delete prompt '{name}'?", parent=self):
            del self.parent.settings['post_editing_prompts'][name]
            save_settings(self.parent.settings)
            self._update_prompt_combo()
            first_name = list(self.parent.settings['post_editing_prompts'].keys())[0]
            self.prompt_var.set(first_name)
            self._on_prompt_select()
    
    def _check_for_resume_task(self):
        if os.path.exists(RESUME_PE_FILE):
            try:
                with open(RESUME_PE_FILE, 'r', encoding='utf-8') as f: data = json.load(f)
                file_name = os.path.basename(data.get('current_file', 'unknown file'))
                row_index = data.get('last_row_index', -1)
                if messagebox.askyesno("Unfinished Task", f"An unfinished post-editing task for '{file_name}' (stopped at row {row_index + 1}) was found.\n\nDo you want to resume?", parent=self):
                    self._load_resume_state(data)
                else:
                    os.remove(RESUME_PE_FILE)
            except Exception as e:
                log_error(f"Failed to read post-edit resume file: {e}")
                if os.path.exists(RESUME_PE_FILE): os.remove(RESUME_PE_FILE)
    
    def _load_resume_state(self, data):
        self.resume_data = data
        self.selected_files = data.get('all_files', [])
        self.file_listbox.delete(0, tk.END)
        for f in self.selected_files:
            self.file_listbox.insert(tk.END, os.path.basename(f))
        self._update_status(f"Ready to resume. {len(self.selected_files)} files loaded.", "blue")
    
    def _save_resume_state(self, current_file, last_index, edited_paras, all_files):
        state = {
            'current_file': current_file,
            'last_row_index': last_index,
            'edited_paragraphs': edited_paras,
            'all_files': all_files
        }
        try:
            with open(RESUME_PE_FILE, 'w', encoding='utf-8') as f: json.dump(state, f, indent=4)
        except Exception as e:
            log_error(f"Failed to save post-edit resume state: {e}")
    
    def _start_post_editing(self):
        if not self.selected_files: return messagebox.showerror("Error", "Please select one or more Excel files.", parent=self)
        prompt = self.prompt_text.get("1.0", tk.END).strip()
        if not prompt: return messagebox.showerror("Error", "Prompt cannot be empty.", parent=self)
        if "{source}" not in prompt or "{target}" not in prompt:
            return messagebox.showerror("Error", "Prompt must contain {source} and {target} placeholders.", parent=self)
    
        self.is_processing = True
        self.stop_requested.clear()
        self.process_button.config(text="Stop Processing", command=self._stop_post_editing)
        self._update_status("Processing...", "orange")
    
        threading.Thread(target=self._post_editing_task, args=(self.resume_data,), daemon=True).start()
        self.resume_data = None
    
    def _stop_post_editing(self):
        self.stop_requested.set()
        self._cancel_timer()
        self._update_status("Stopping...", "orange")
    
    def _post_editing_task(self, resume_data=None):
        try:
            model_name = self.parent.model_name_var.get().strip()
            max_tokens = self.parent.settings.get('max_tokens', 8000)
            retry_attempts = self.parent.settings.get('retry_attempts', 3)
            paragraph_timeout = self.parent.settings.get('paragraph_timeout', 300)
            request_interval_value = self.parent.settings.get('request_interval', 5)
            prompt_template = self.prompt_text.get("1.0", tk.END).strip()
            output_format = self.parent.settings.get('output_format', 'Both')
            output_location = self.parent.settings.get('output_location', 'Subfolder')
            
            client = self.parent._create_client()
            
            total_files = len(self.selected_files)
            start_file_index = 0
            if resume_data:
                try:
                    start_file_index = self.selected_files.index(resume_data.get('current_file'))
                except ValueError:
                    log_error(f"Resumed file '{resume_data.get('current_file')}' not found in selection.")
    
            for file_idx in range(start_file_index, total_files):
                file_path = self.selected_files[file_idx]
                file_name = os.path.basename(file_path)
                
                base_name = os.path.splitext(file_name)[0].replace('_corpus', '').replace('_translated','')
    
                if output_location == "Subfolder":
                    output_dir = os.path.join(os.path.dirname(file_path), base_name)
                    os.makedirs(output_dir, exist_ok=True)
                else:
                    output_dir = os.path.dirname(file_path)
    
                self.after(0, self._update_status, f"[{file_idx+1}/{total_files}] Reading: {file_name}", "orange")
                df = pd.read_excel(file_path)
                
                if 'Source' not in df.columns or 'Translation' not in df.columns:
                    log_error(f"File {file_name} skipped: must contain 'Source' and 'Translation' columns.")
                    continue
    
                edited_paragraphs = []
                total_rows = len(df)
                start_row = 0
                if resume_data and file_path == resume_data.get('current_file'):
                    start_row = resume_data.get('last_row_index', -1) + 1
                    edited_paragraphs = resume_data.get('edited_paragraphs', [])
                    resume_data = None
    
                for i, row in df.iloc[start_row:].iterrows():
                    if self.stop_requested.is_set():
                        self._save_resume_state(file_path, i - 1, edited_paragraphs, self.selected_files)
                        self.after(0, self._update_status, f"Stopped. Progress for '{file_name}' saved.", "blue")
                        return
    
                    self.after(0, self._update_status, f"[{file_idx+1}/{total_files}] Editing {file_name} (row {i + 1}/{total_rows})", "orange")
                    start_time = time.time()
                    self.after(0, self._update_timer, start_time)
    
                    source_text, target_text = str(row['Source']), str(row['Translation'])
                    full_prompt = prompt_template.format(source=source_text, target=target_text)
                    
                    edited_para = translate_single_paragraph(client, model_name, full_prompt, max_tokens, retry_attempts, paragraph_timeout)
                    
                    self.after(0, self._cancel_timer)
                    edited_paragraphs.append(edited_para)
                    
                    if request_interval_value > 0 and i < total_rows - 1:
                        time.sleep(request_interval_value)
                
                self.after(0, self._update_status, f"[{file_idx+1}/{total_files}] Saving output for {file_name}...", "orange")
                
                df['Post-edited'] = pd.Series(edited_paragraphs)
                
                if output_format in ["Excel", "Both"]:
                    excel_path = os.path.join(output_dir, f"{base_name}_postedited.xlsx")
                    df.to_excel(excel_path, index=False, engine='openpyxl')
                    
                    red_bold_font = Font(color="FF0000", bold=True)
                    wb = load_workbook(excel_path)
                    ws = wb.active
                    post_edited_col_idx = -1
                    for col_idx, cell in enumerate(ws[1]):
                        if cell.value == 'Post-edited':
                            post_edited_col_idx = col_idx
                            break
                    
                    if post_edited_col_idx != -1:
                        for row_ws in ws.iter_rows(min_row=2, max_row=ws.max_row):
                            cell = row_ws[post_edited_col_idx]
                            if not isinstance(cell.value, str): continue
                            if cell.value == "[ERROR_CONTENT_FILTER]":
                                cell.value = "Rejected by API (content policy)"
                                cell.font = red_bold_font
                            elif cell.value == "[ERROR_NETWORK]":
                                cell.value = "Network Issue"
                                cell.font = red_bold_font
                            elif cell.value.startswith("[ERROR_OTHER:"):
                                cell.value = f"Failed: {cell.value[13:-3]}"
                                cell.font = red_bold_font
                    wb.save(excel_path)
                
                if output_format in ["Text", "Both"]:
                    if 'Post-edited' in df.columns and not df['Post-edited'].isnull().all():
                        full_edited_text = "\n\n".join(df['Post-edited'].astype(str).tolist())
                        txt_path = os.path.join(output_dir, f"{base_name}_postedited.txt")
                        with open(txt_path, 'w', encoding='utf-8') as f: f.write(full_edited_text)
            
            if os.path.exists(RESUME_PE_FILE): os.remove(RESUME_PE_FILE)
            self.after(0, self._update_status, "Post-editing complete! All files saved.", "green")
    
        except Exception as e:
            error_message = f"Processing failed: {e}"
            log_error(f"Post-editing task failed. Error: {e}")
            self.after(0, self._update_status, error_message, "red")
            self.after(0, messagebox.showerror, "An Error Occurred", f"{e}\n\nDetails logged to error_log.txt", parent=self)
        
        finally:
            self.is_processing = False
            self.after(0, lambda: self.process_button.config(text="Start Post-editing", command=self._start_post_editing))
            self.after(0, self._cancel_timer)